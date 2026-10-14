# Membuat content/gending-isian.xlsx (format isian gending untuk guru).
#   python scripts/make-gending-template.py
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

F = 'Arial'
INK = '2A1C12'; MUTED = '7A6554'; ACCENT = 'A8431F'
INPUT = PatternFill('solid', fgColor='FFF4C7')   # kolom yang diisi guru
AUTO = PatternFill('solid', fgColor='EFE8DD')    # terisi otomatis
HEAD = PatternFill('solid', fgColor='4A2614')
thin = Side(style='thin', color='D9CBB5')
box = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---------------- Petunjuk ----------------
p = wb.active
p.title = 'Petunjuk'
p.sheet_view.showGridLines = False
p.column_dimensions['A'].width = 3
p.column_dimensions['B'].width = 26
p.column_dimensions['C'].width = 80
r = 2


def put(text, size=10, bold=False, color=INK, col='B', height=None):
    global r
    c = p[f'{col}{r}']
    c.value = text
    c.font = Font(name=F, size=size, bold=bold, color=color)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if height:
        p.row_dimensions[r].height = height
    r += 1


put('Nadantara — Isian Gending', 16, True, height=24)
p.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
put('Untuk guru karawitan: tulis gending di lembar "Gending", satu gending per baris. '
    'File ini lalu dimasukkan ke aplikasi Nadantara.', 10, color=MUTED, height=28)
r += 1
put('Cara mengisi', 12, True, ACCENT, height=20)
steps = [
    'Buka lembar "Gending". Kolom KUNING diisi, kolom ABU-ABU terisi otomatis. Tanda * berarti wajib.',
    'Pilih Bentuk, Laras, Pathet, dan Jenjang dari daftar (klik sel, lalu panah kecil di kanan sel).',
    'Tulis balungan dengan aturan di bawah. Periksa kolom "Pemeriksaan": harus tertulis "OK".',
    'Baris 5 adalah CONTOH latihan, bukan gending asli. Biarkan saja; baris itu tidak ikut dimasukkan.',
    'Kirim file yang sudah diisi. Gending akan muncul di menu "Contoh latihan" di aplikasi.',
]
for i, s in enumerate(steps, 1):
    c = p[f'B{r}']
    c.value = f'{i}.'
    c.font = Font(name=F, size=10, bold=True, color=INK)
    c.alignment = Alignment(horizontal='right', vertical='top')
    put(s, col='C', height=16)
r += 1
put('Aturan menulis balungan', 12, True, ACCENT, height=20)
rows = [
    ('Tulis', 'Arti'),
    ('1  2  3  4  5  6  7', 'Nada. Sléndro hanya memakai 1 2 3 5 6.'),
    ('.', 'Pin: ketukan yang tidak dipukul.'),
    ('6,   (angka + koma)', 'Nada rendah = angka dengan titik di BAWAH. Contoh: 6,  5,'),
    ("1'   (angka + petik)", "Nada tinggi = angka dengan titik di ATAS. Contoh: 1'  2'"),
    ('spasi', 'Bebas, hanya untuk kerapian. Biasanya satu spasi di antara gatra: .5.3 .5.3'),
    ('(2)  dan tanda lain', 'Tanda gong/kenong dalam kurung, huruf, dan garis diabaikan. '
                           'Gong dan kenong dihitung otomatis dari bentuk gending.'),
]
for i, (a, b) in enumerate(rows):
    ca, cb = p[f'B{r}'], p[f'C{r}']
    ca.value, cb.value = a, b
    head = i == 0
    for c in (ca, cb):
        c.font = Font(name=F, size=10, bold=head, color='FFFFFF' if head else INK)
        c.fill = HEAD if head else PatternFill()
        c.border = box
        c.alignment = Alignment(wrap_text=True, vertical='center')
    if not head:
        ca.font = Font(name='Consolas', size=11, bold=True, color=INK)
    p.row_dimensions[r].height = 30 if len(b) > 85 else 18
    r += 1
r += 1
put('Contoh penulisan', 12, True, ACCENT, height=20)
examples = [
    ('Di buku (kepatihan)', '. 5 . 3    . 5 . 3    . 6 . 5    . 3 . (2)'),
    ('Ditulis di sini', '.5.3 .5.3 .6.5 .3.2'),
    ('Nada rendah & tinggi', "6, 1 2 3   5 6 1' .      (6 rendah … 1 tinggi)"),
]
for a, b in examples:
    p[f'B{r}'] = a
    p[f'B{r}'].font = Font(name=F, size=10, color=MUTED)
    p[f'C{r}'] = b
    p[f'C{r}'].font = Font(name='Consolas', size=11, color=INK)
    r += 1
r += 1
put('Bentuk yang sudah bisa diputar', 12, True, ACCENT, height=20)
put('Lancaran (16 ketukan per gongan) dan Bebas. Ketawang, Ladrang, dan Gangsaran boleh diisi sekarang; '
    'untuk sementara diputar sebagai "Bebas" sampai polanya didukung aplikasi.', col='C', height=30)

# ---------------- Gending ----------------
g = wb.create_sheet('Gending')
g.sheet_view.showGridLines = False
cols = [
    ('No', 7, 'auto'), ('Judul *', 32, 'in'), ('Bentuk *', 13, 'in'), ('Laras *', 11, 'in'),
    ('Pathet', 11, 'in'), ('Jenjang', 10, 'in'), ('Buka', 22, 'in'), ('Balungan *', 58, 'in'),
    ('Jumlah ketukan', 11, 'auto'), ('Pemeriksaan', 36, 'auto'),
    ('Sumber / referensi', 26, 'in'), ('Catatan', 28, 'in'), ('Diisi oleh', 16, 'in'),
]
g['A1'] = 'Daftar Gending — Nadantara'
g['A1'].font = Font(name=F, size=14, bold=True, color=INK)
g['A2'] = ('Kolom kuning diisi · kolom abu-abu terisi otomatis · * wajib · '
           'baris 5 = contoh latihan, tidak ikut dimasukkan')
g['A2'].font = Font(name=F, size=9, italic=True, color=MUTED)
HR = 4
for i, (name, w, kind) in enumerate(cols, 1):
    c = g.cell(row=HR, column=i, value=name)
    c.font = Font(name=F, size=10, bold=True, color='FFFFFF')
    c.fill = HEAD
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = box
    g.column_dimensions[c.column_letter].width = w
g.row_dimensions[HR].height = 30
g.freeze_panes = 'C5'

FIRST, LAST = 5, 104   # baris contoh + 99 baris isian
for rr in range(FIRST, LAST + 1):
    g.row_dimensions[rr].height = 20
    for i, (name, w, kind) in enumerate(cols, 1):
        c = g.cell(row=rr, column=i)
        c.border = box
        c.fill = INPUT if kind == 'in' else AUTO
        c.font = Font(name='Consolas' if name in ('Balungan *', 'Buka') else F, size=10, color=INK)
        c.alignment = Alignment(vertical='center', wrap_text=name in ('Balungan *', 'Pemeriksaan', 'Catatan'))
    g[f'A{rr}'] = f'=IF(B{rr}="","",ROW()-{FIRST})'
    g[f'I{rr}'] = (f'=IF(H{rr}="","",SUMPRODUCT(LEN(H{rr})-LEN(SUBSTITUTE(H{rr},'
                   '{"1";"2";"3";"4";"5";"6";"7";"."},""))))')
    g[f'J{rr}'] = (f'=IF(H{rr}="",IF(B{rr}="","","Balungan belum diisi"),'
                   f'IF(OR(C{rr}="",D{rr}=""),"Pilih Bentuk dan Laras",'
                   f'IF(AND(D{rr}="Sléndro",OR(ISNUMBER(FIND("4",H{rr})),ISNUMBER(FIND("7",H{rr})))),'
                   f'"Nada 4/7 tidak ada di sléndro",'
                   f'IF(AND(C{rr}="Lancaran",MOD(I{rr},16)<>0),'
                   f'"Lancaran harus kelipatan 16 ketukan (sekarang "&I{rr}&")","OK"))))')

example = {
    2: 'CONTOH — Latihan pola lancaran (bukan gending asli)', 3: 'Lancaran', 4: 'Sléndro',
    5: 'Manyura', 6: 'SD', 8: '.5.3 .5.3 .6.5 .3.2',
    11: 'Contoh dari aplikasi Nadantara', 12: 'Baris contoh; tidak ikut dimasukkan', 13: 'Nadantara',
}
for col, v in example.items():
    g.cell(row=FIRST, column=col, value=v)
g.cell(row=FIRST, column=1, value='contoh')
for i in range(1, len(cols) + 1):
    c = g.cell(row=FIRST, column=i)
    c.font = Font(name=c.font.name, size=10, italic=True, color=MUTED)
    c.fill = PatternFill('solid', fgColor='F3EEE6')


def dv(formula, col, prompt):
    d = DataValidation(type='list', formula1=formula, allow_blank=True)
    d.promptTitle, d.prompt, d.showInputMessage = 'Pilih dari daftar', prompt, True
    d.errorTitle, d.error, d.showErrorMessage = 'Pilihan tidak dikenal', 'Pilih salah satu dari daftar.', True
    g.add_data_validation(d)
    d.add(f'{col}{FIRST}:{col}{LAST}')


dv('"Lancaran,Ketawang,Ladrang,Gangsaran,Bebas"', 'C', 'Lancaran dan Bebas sudah bisa diputar.')
dv('"Sléndro,Pélog"', 'D', 'Sléndro: 1 2 3 5 6. Pélog: 1 sampai 7.')
dv('"Nem,Sanga,Manyura,Lima,Barang"', 'E', 'Sléndro: Nem/Sanga/Manyura. Pélog: Lima/Nem/Barang.')
dv('"SD,SMP,SMA,Semua"', 'F', 'Untuk jenjang kelas mana.')

# petunjuk singkat di atas judul kolom (selalu terlihat, tanpa perlu diarahkan kursor)
hints = {'G3': 'opsional, aturan sama', 'H3': "6, = rendah · 1' = tinggi · . = pin · lihat Petunjuk"}
for ref, t in hints.items():
    g[ref] = t
    g[ref].font = Font(name=F, size=8, italic=True, color=MUTED)
    g[ref].alignment = Alignment(vertical='bottom')
g.auto_filter.ref = f'A{HR}:M{LAST}'
wb.active = 1
wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'content', 'gending-isian.xlsx'))
print('ok')
